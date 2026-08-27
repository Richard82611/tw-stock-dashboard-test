# -*- coding: utf-8 -*-
"""
全市場三率最終顯示修正 v2.2
============================

背景：
v2.1 底層 _income_validation_df 已確認 1953 檔中：
- 1911 檔「可計算三率」
- 34 檔「三率不適用」
- 真正未解決只剩 8 檔

但 fundamental_full_market_fix.py 的 prof_map 主要由「最新 snapshot」建立，
少數股票雖然 SQLite 歷史季報已完整、驗收也是「可計算三率」，
卻沒有進 prof_map，導致 Excel / JSON 被錯列成資料缺漏。

本程式不重抓資料、不改三率定義。
只用 SQLite 中 target 日以前已存在且已驗證的官方歷史季報，
重新計算並覆寫 duck_latest.xlsx 的三率顯示與 fundamental_full_market_status.json。

三率固定：
- 毛利率
- 營業利益率
- 稅後淨利率
最新單季 vs 前一單季，三項同時上升（依 settings.ini allow_equal）才是三率三升。
"""
from __future__ import annotations

import argparse
import datetime as dt
import json
import math
import sqlite3
import sys
from pathlib import Path
from typing import Dict, Optional, Tuple

import pandas as pd
from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
sys.path.insert(0, str(HERE))

import stock_duckbill as sd  # noqa: E402


RATE_FIELDS = [
    "毛利率本季%","毛利率上季%",
    "營業利益率本季%","營業利益率上季%",
    "稅後純益率本季%","稅後純益率上季%",
    "毛差","營差","淨差",
    "三率資料期","三率快照日","三率可得日依據","三率來源",
]


def target_date(s: Optional[str]) -> dt.date:
    if s:
        return dt.datetime.strptime(s, "%Y-%m-%d").date()
    now = dt.datetime.now(dt.timezone(dt.timedelta(hours=8)))
    d = now.date() if now.hour >= 16 else now.date() - dt.timedelta(days=1)
    while d.weekday() >= 5:
        d -= dt.timedelta(days=1)
    return d


def _prev_period(y: int, q: int) -> Tuple[int, int]:
    return (y, q - 1) if q > 1 else (y - 1, 4)


def _single(pmap: Dict[Tuple[int,int], dict], y: int, q: int) -> Optional[dict]:
    cur = pmap.get((y, q))
    if not cur:
        return None
    core = ["revenue_ytd","gross_profit_ytd","operating_income_ytd","net_income_ytd"]
    if any(cur.get(k) is None or pd.isna(cur.get(k)) for k in core):
        return None
    if q == 1:
        return {k: float(cur[k]) for k in core}
    prev = pmap.get((y, q - 1))
    if not prev or any(prev.get(k) is None or pd.isna(prev.get(k)) for k in core):
        return None
    return {k: float(cur[k]) - float(prev[k]) for k in core}


def calc_three_rates_asof(
    conn: sqlite3.Connection,
    sid: str,
    target: dt.date,
    allow_equal: bool,
) -> Optional[dict]:
    rows = pd.read_sql_query(
        """SELECT year,quarter,revenue_ytd,gross_profit_ytd,
                  operating_income_ytd,net_income_ytd,eps_ytd,
                  availability_date,availability_basis,source
           FROM fundamental_income_history
           WHERE stock_id=? AND availability_date IS NOT NULL AND availability_date<=?
           ORDER BY year,quarter""",
        conn,
        params=(str(sid), target.isoformat()),
    )
    if rows.empty:
        return None

    pmap = {}
    for r in rows.to_dict("records"):
        pmap[(int(r["year"]), int(r["quarter"]))] = r

    cy, cq = max(pmap.keys())
    py, pq = _prev_period(cy, cq)
    cur = _single(pmap, cy, cq)
    prev = _single(pmap, py, pq)
    if not cur or not prev:
        return None
    if abs(cur["revenue_ytd"]) < 1e-12 or abs(prev["revenue_ytd"]) < 1e-12:
        return None

    cg = cur["gross_profit_ytd"] / cur["revenue_ytd"]
    pg = prev["gross_profit_ytd"] / prev["revenue_ytd"]
    co = cur["operating_income_ytd"] / cur["revenue_ytd"]
    po = prev["operating_income_ytd"] / prev["revenue_ytd"]
    cn = cur["net_income_ytd"] / cur["revenue_ytd"]
    pn = prev["net_income_ytd"] / prev["revenue_ytd"]
    gd, od, nd = cg-pg, co-po, cn-pn
    cmp = (lambda x: x >= -1e-12) if allow_equal else (lambda x: x > 1e-12)

    latest = pmap[(cy, cq)]
    return {
        "三率三升": bool(cmp(gd) and cmp(od) and cmp(nd)),
        "毛利率本季%": cg*100,
        "毛利率上季%": pg*100,
        "營業利益率本季%": co*100,
        "營業利益率上季%": po*100,
        "稅後純益率本季%": cn*100,
        "稅後純益率上季%": pn*100,
        "毛差": gd,
        "營差": od,
        "淨差": nd,
        "三率資料期": f"{cy}Q{cq} vs {py}Q{pq}",
        "三率快照日": str(latest.get("availability_date") or target.isoformat()),
        "三率可得日依據": str(latest.get("availability_basis") or "sqlite_official_history"),
        "三率來源": str(latest.get("source") or "官方季報") + "＋SQLite官方歷史季報",
    }


def _headers(ws, row=2):
    return {
        str(ws.cell(row=row, column=c).value or "").strip(): c
        for c in range(1, ws.max_column + 1)
    }


def _ensure_cols(ws, headers, names):
    for name in names:
        if name not in headers:
            c = ws.max_column + 1
            ws.cell(row=2, column=c, value=name)
            headers[name] = c


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--date")
    ap.add_argument("--workbook", default=str(ROOT / "duck_latest.xlsx"))
    args = ap.parse_args()

    target = target_date(args.date)
    workbook = Path(args.workbook)
    if not workbook.is_absolute():
        workbook = (Path.cwd() / workbook).resolve()
    if not workbook.exists():
        raise FileNotFoundError(workbook)

    cfg = sd.load_config()
    allow_equal = cfg.getboolean("general", "three_rates_allow_equal", fallback=True)

    conn = sqlite3.connect(sd.DB_PATH)
    sd.init_db(conn)

    wb = load_workbook(workbook)
    if "全市場整合篩選" not in wb.sheetnames:
        raise RuntimeError("找不到『全市場整合篩選』")
    ws = wb["全市場整合篩選"]
    h = _headers(ws)
    if "代號" not in h:
        raise RuntimeError("找不到『代號』欄")

    _ensure_cols(
        ws, h,
        ["三率狀態","基本面資料狀態","基本面缺漏原因","基本面所需季度"] + RATE_FIELDS
    )

    stock_rows = []
    for r in range(3, ws.max_row + 1):
        sidv = ws.cell(r, h["代號"]).value
        if sidv is None:
            continue
        sid = str(sidv).replace(".0","").strip()
        if sid:
            stock_rows.append((r, sid))

    ids = [sid for _, sid in stock_rows]
    vdf = sd._income_validation_df(conn, target, ids)
    vmap = {
        str(x["stock_id"]): x
        for _, x in vdf.iterrows()
    } if vdf is not None and not vdf.empty else {}

    counts = {
        "stocks": 0,
        "three_rate_yes": 0,
        "three_rate_no": 0,
        "three_rate_na": 0,
        "three_rate_listing_insufficient": 0,
        "three_rate_prior_history_short": 0,
        "three_rate_data_missing": 0,
        "three_rate_unknown": 0,
    }
    audit = []
    quanta = None

    for row, sid in stock_rows:
        counts["stocks"] += 1
        vv = vmap.get(sid)
        bs = str(vv.get("status") or "") if vv is not None else ""
        reason = str(vv.get("reason") or "") if vv is not None else ""
        required = str(vv.get("required") or "") if vv is not None else ""

        # 先清舊三率詳細欄，避免殘留。
        for k in RATE_FIELDS:
            ws.cell(row, h[k], value=None)

        result = None
        if bs == "可計算三率":
            result = calc_three_rates_asof(conn, sid, target, allow_equal)
            if result:
                status = "三率三升" if result["三率三升"] else "未三升"
                for k in RATE_FIELDS:
                    ws.cell(row, h[k], value=result.get(k))
                if result["三率三升"]:
                    counts["three_rate_yes"] += 1
                else:
                    counts["three_rate_no"] += 1
            else:
                # 驗收說可算但本程式仍算不出，才是真正需要警示的異常。
                status = "資料缺漏"
                reason = "底層驗收為可計算三率，但最終重算失敗，請檢查季度欄位"
                counts["three_rate_data_missing"] += 1
        elif bs == "三率不適用":
            status = "三率不適用"
            counts["three_rate_na"] += 1
        elif bs == "上市時間不足/不適用":
            status = "上市時間不足"
            counts["three_rate_listing_insufficient"] += 1
        elif bs == "比較期歷史不足":
            status = "比較期歷史不足"
            counts["three_rate_prior_history_short"] += 1
        elif bs == "缺資料":
            status = "資料缺漏"
            counts["three_rate_data_missing"] += 1
        else:
            status = "資料不足"
            counts["three_rate_unknown"] += 1

        ws.cell(row, h["三率狀態"], value=status)
        ws.cell(row, h["基本面資料狀態"], value=bs or status)
        ws.cell(row, h["基本面缺漏原因"], value=reason)
        ws.cell(row, h["基本面所需季度"], value=required)

        name = ""
        if "名稱" in h:
            name = str(ws.cell(row, h["名稱"]).value or "")
        revenue_status = str(ws.cell(row, h.get("營收狀態", 0)).value or "") if h.get("營收狀態") else ""

        audit.append([sid, name, status, result.get("三率資料期") if result else "", revenue_status, reason])

        if sid == "2382":
            quanta = {
                "stock_id": sid,
                "name": name,
                "three_rate_status": status,
                "three_rate_period": result.get("三率資料期") if result else None,
                "gross_current_pct": result.get("毛利率本季%") if result else None,
                "gross_prev_pct": result.get("毛利率上季%") if result else None,
                "operating_current_pct": result.get("營業利益率本季%") if result else None,
                "operating_prev_pct": result.get("營業利益率上季%") if result else None,
                "net_current_pct": result.get("稅後純益率本季%") if result else None,
                "net_prev_pct": result.get("稅後純益率上季%") if result else None,
            }

    # 重建驗收表，避免舊的「79檔資料不足」留在檔案裡。
    if "基本面全市場驗收" in wb.sheetnames:
        del wb["基本面全市場驗收"]
    va = wb.create_sheet("基本面全市場驗收")
    va.append(["代號","名稱","三率狀態","三率資料期","營收狀態","缺漏/說明"])
    for x in audit:
        va.append(x)
    va.freeze_panes = "A2"
    va.auto_filter.ref = va.dimensions
    for idx, width in enumerate([12,18,22,20,22,60], 1):
        va.column_dimensions[chr(64+idx)].width = width

    va["H1"] = "最終驗收摘要 v2.2"
    summary_lines = [
        f"目標日：{target}",
        f"普通股：{counts['stocks']}",
        f"三率三升：{counts['three_rate_yes']}",
        f"未三升：{counts['three_rate_no']}",
        f"三率不適用：{counts['three_rate_na']}",
        f"上市時間不足：{counts['three_rate_listing_insufficient']}",
        f"較舊比較期不足：{counts['three_rate_prior_history_short']}",
        f"真正資料缺漏：{counts['three_rate_data_missing']}",
        f"其他未分類：{counts['three_rate_unknown']}",
    ]
    for i, v in enumerate(summary_lines, 2):
        va.cell(i, 8, v)
    va.column_dimensions["H"].width = 38

    wb.save(workbook)

    # 保留原 report 的營收統計/來源，只覆蓋三率統計。
    report_path = ROOT / "fundamental_full_market_status.json"
    try:
        old = json.loads(report_path.read_text(encoding="utf-8"))
    except Exception:
        old = {}

    old_counts = dict(old.get("counts") or {})
    old_counts.update(counts)
    old["target_date"] = target.isoformat()
    old["workbook"] = str(workbook)
    old["all_market_stock_count"] = counts["stocks"]
    old["counts"] = old_counts
    old["quanta_2382"] = {
        **(old.get("quanta_2382") or {}),
        **(quanta or {}),
    }
    old["three_rate_finalizer"] = {
        "version": "v2.2",
        "source": "SQLite official history as-of target",
        "validation_source": "_income_validation_df",
        "snapshot_only_false_missing_fixed": True,
        "allow_equal": allow_equal,
    }
    report_path.write_text(json.dumps(old, ensure_ascii=False, indent=2), encoding="utf-8")

    status_path = ROOT / "fundamental_three_rate_finalizer_status.json"
    status = {
        "version": "three-rate-finalizer-v2.2",
        "target_date": target.isoformat(),
        "counts": counts,
        "quanta_2382": quanta,
        "note": "以底層驗收＋SQLite官方歷史季報重建Excel三率顯示，修正snapshot-only造成的假缺漏。",
    }
    status_path.write_text(json.dumps(status, ensure_ascii=False, indent=2), encoding="utf-8")

    print("=" * 72)
    print("[v2.2 最終顯示修正完成]")
    print(json.dumps(counts, ensure_ascii=False, indent=2))
    print("2382:", json.dumps(quanta, ensure_ascii=False, indent=2))
    print("=" * 72)
    conn.close()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
