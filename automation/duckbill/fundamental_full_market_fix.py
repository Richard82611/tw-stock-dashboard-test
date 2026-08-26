# -*- coding: utf-8 -*-
"""
全市場基本面驗收／修復器 v1.0
================================
目的：
1. 三率三升改為「全市場普通股」都檢查，不再只限正式鴨嘴＋預備鴨嘴。
2. 營收三增改為「全市場普通股」都驗收。
3. 「沒有資料／基期不足」不得誤標成 False / 未三增 / 未三升。
4. 第一次可用 --repair-missing-core 只補三率判定真正需要的核心季度；
   不需要每檔補12季，日常更新不做深度歷史重抓。
5. 直接修正 duck_latest.xlsx 的「全市場整合篩選」，不改鴨嘴技術面與估值公式。

三率定義固定：
- 毛利率
- 營業利益率
- 稅後淨利率
最新單季 vs 前一單季，三者都上升才是「三率三升」。
"""
from __future__ import annotations

import argparse
import datetime as dt
import json
import math
import sqlite3
import sys
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from pathlib import Path
from typing import Dict, List, Optional, Tuple

import pandas as pd
from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
sys.path.insert(0, str(HERE))

import stock_duckbill as sd  # noqa: E402


def _num(v):
    try:
        if v is None or pd.isna(v):
            return None
        x = float(v)
        return x if math.isfinite(x) else None
    except Exception:
        return None


def _target_date(s: Optional[str]) -> dt.date:
    if s:
        return dt.datetime.strptime(s, "%Y-%m-%d").date()
    now = dt.datetime.now(dt.timezone(dt.timedelta(hours=8)))
    day = now.date() if now.hour >= 16 else now.date() - dt.timedelta(days=1)
    while day.weekday() >= 5:
        day -= dt.timedelta(days=1)
    return day


def _all_market_ids(conn: sqlite3.Connection, target: dt.date) -> List[str]:
    rows = conn.execute(
        "SELECT stock_id,name FROM prices WHERE date=? ORDER BY stock_id",
        (target.isoformat(),),
    ).fetchall()
    out = []
    for sid, name in rows:
        sid = str(sid)
        if sd.is_common_stock(sid, str(name or "")):
            out.append(sid)
    return list(dict.fromkeys(out))


def _existing_income_periods(conn, sid: str, target: dt.date) -> set[Tuple[int, int]]:
    rows = conn.execute(
        """SELECT year,quarter
           FROM fundamental_income_history
           WHERE stock_id=? AND availability_date IS NOT NULL AND availability_date<=?""",
        (sid, target.isoformat()),
    ).fetchall()
    return {(int(y), int(q)) for y, q in rows}


def _repair_one_income(sid: str, y: int, q: int):
    """只做 HTTP 取得；SQLite 寫入留在主執行緒，避免多執行緒資料庫鎖。"""
    try:
        r = sd._fetch_mops_income_history_one(sid, y, q)
        return sid, y, q, r, None
    except Exception as e:
        return sid, y, q, None, f"{type(e).__name__}: {e}"


def repair_missing_core_income(
    conn: sqlite3.Connection,
    target: dt.date,
    ids: List[str],
    workers: int = 4,
) -> dict:
    """
    只補「今天判斷三率真正需要」的季度。
    例如 2026Q2 vs Q1 通常只需要 Q1/Q2；
    Q3 則需 Q1/Q2/Q3 來還原單季。
    """
    tasks: List[Tuple[str, int, int]] = []
    skipped_na = 0
    cached_nodata = 0

    for sid in ids:
        if sd._is_three_rate_na_stock(conn, sid, target):
            skipped_na += 1
            continue
        req = sd._three_rate_required_periods_for_stock(conn, sid, target)
        have = _existing_income_periods(conn, sid, target)
        for y, q in req:
            if (y, q) in have:
                continue
            period = f"{y}Q{q}"
            cached = conn.execute(
                """SELECT 1 FROM fundamental_backfill_nodata
                   WHERE dataset='income' AND stock_id=? AND period=?""",
                (sid, period),
            ).fetchone()
            if cached:
                cached_nodata += 1
                continue
            tasks.append((sid, y, q))

    print(
        f"[全市場三率核心補洞] 普通股 {len(ids)} 檔；"
        f"需抓 {len(tasks)} 個核心季度；三率不適用 {skipped_na} 檔；"
        f"已知查無 {cached_nodata} 筆。",
        flush=True,
    )

    added = nodata = failed = 0
    errors = []
    if not tasks:
        return {
            "tasks": 0, "added": 0, "nodata": 0, "failed": 0,
            "skipped_na": skipped_na, "cached_nodata": cached_nodata,
        }

    workers = max(1, min(int(workers), 6))
    with ThreadPoolExecutor(max_workers=workers) as ex:
        futs = [ex.submit(_repair_one_income, *t) for t in tasks]
        for n, fut in enumerate(as_completed(futs), 1):
            sid, y, q, rec, err = fut.result()
            period = f"{y}Q{q}"
            if rec:
                try:
                    sd._upsert_income_history_record(conn, rec)
                    added += 1
                except Exception as e:
                    failed += 1
                    errors.append(f"{sid}:{period} DB:{e}")
            elif err:
                failed += 1
                errors.append(f"{sid}:{period} {err}")
            else:
                # stock_duckbill 的舊頁解析器只有在官方明確查無時才回 None；
                # 解析異常會 raise，因此這裡可以安全寫 NODATA。
                conn.execute(
                    """INSERT OR REPLACE INTO fundamental_backfill_nodata
                       (dataset,stock_id,period,checked_at)
                       VALUES('income',?,?,?)""",
                    (sid, period, dt.datetime.now().isoformat(timespec="seconds")),
                )
                conn.commit()
                nodata += 1
            if n == 1 or n % 50 == 0 or n == len(tasks):
                print(
                    f"[核心補洞 {n}/{len(tasks)}] 新增={added} 查無={nodata} 失敗={failed}",
                    flush=True,
                )

    return {
        "tasks": len(tasks), "added": added, "nodata": nodata,
        "failed": failed, "skipped_na": skipped_na,
        "cached_nodata": cached_nodata, "errors": errors[:50],
    }


def build_full_market_maps(conn: sqlite3.Connection, target: dt.date, ids: List[str], cfg):
    # 先修正特殊產業/KY的保守可得日期。
    sd._normalize_special_availability_dates(conn, target)

    # 最新季報：TWSE/TPEx 六類官方批次端點，以「全市場」ID進行驗收/保存。
    bulk_added, bulk_errors = sd._seed_bulk_income_for_target(conn, target, ids)

    # 月營收：官方批次資料本身是全市場；全市場ID讓增量完整性檢查也不再只看鴨嘴候選。
    revenue = sd.fetch_official_month_revenue(conn, target, cfg, ids)
    recalc, recalc_na = sd._recompute_revenue_growth_from_history(conn, target, ids)
    if recalc:
        revenue = sd._latest_revenue_history_asof(conn, target, None)

    # 最新財報 snapshot / history。
    income_snap = sd.fetch_official_income_snapshot(conn, target, ids)
    allow_equal = cfg.getboolean("general", "three_rates_allow_equal", fallback=True)
    revmin = cfg.getfloat("general", "revenue_growth_min_pct", fallback=0.0)

    prof_map: Dict[str, dict] = {}
    if income_snap is not None and not income_snap.empty:
        sub = income_snap[income_snap["stock_id"].astype(str).isin(set(ids))]
        for _, snap in sub.iterrows():
            sid = str(snap["stock_id"])
            x = sd._three_rates_for_stock(conn, snap, allow_equal)
            if x:
                prof_map[sid] = x

    # 全市場營收：None 代表不是「否」，而是無法可靠計算。
    rev_map: Dict[str, dict] = {}
    rev_raw_map: Dict[str, dict] = {}
    rev_all = sd._latest_revenue_history_asof(conn, target, None)
    if rev_all is not None and not rev_all.empty:
        for _, r in rev_all.iterrows():
            sid = str(r["stock_id"])
            if sid not in set(ids):
                continue
            mk = int(r["month_key"]) if pd.notna(r.get("month_key")) else None
            mom = sd.clean_num(r.get("revenue_mom_pct"))
            yoy = sd.clean_num(r.get("revenue_yoy_pct"))
            cum = sd.clean_num(r.get("revenue_cum_yoy_pct"))
            raw = {
                "營收月增%": mom, "營收年增%": yoy, "累計營收增%": cum,
                "營收資料月": f"{mk//100}-{mk%100:02d}" if mk else None,
                "營收快照日": str(r.get("observed_date") or ""),
                "營收可得日依據": str(r.get("availability_basis") or ""),
                "營收來源": str(r.get("source") or ""),
            }
            rev_raw_map[sid] = raw
            if None not in (mom, yoy, cum):
                raw = dict(raw)
                raw["營收三增"] = bool(
                    mom > revmin and yoy > revmin and cum > revmin
                )
                rev_map[sid] = raw

    vdf = sd._income_validation_df(conn, target, ids)
    vmap = (
        {str(r["stock_id"]): r for _, r in vdf.iterrows()}
        if vdf is not None and not vdf.empty else {}
    )
    return prof_map, rev_map, rev_raw_map, vmap, {
        "bulk_income_added": bulk_added,
        "bulk_income_errors": bulk_errors[:20],
        "revenue_recalc": recalc,
        "revenue_recalc_na": recalc_na,
    }


def three_rate_status(sid: str, prof_map: dict, vmap: dict) -> str:
    vv = vmap.get(sid)
    bs = str(vv.get("status") or "") if vv is not None else ""
    if bs == "三率不適用":
        return "三率不適用"
    if bs == "上市時間不足/不適用":
        return "上市時間不足"
    if sid in prof_map:
        return "三率三升" if bool(prof_map[sid].get("三率三升")) else "未三升"
    if bs == "缺資料":
        return "資料缺漏"
    if bs == "比較期歷史不足":
        return "比較期歷史不足"
    return "資料不足"


def revenue_status(sid: str, rev_map: dict, rev_raw_map: dict) -> str:
    if sid in rev_map:
        return "營收三增" if bool(rev_map[sid].get("營收三增")) else "未三增"
    if sid in rev_raw_map:
        return "成長率不適用/基期不足"
    return "資料不足"


def patch_workbook(
    workbook: Path,
    prof_map: dict,
    rev_map: dict,
    rev_raw_map: dict,
    vmap: dict,
    target: dt.date,
):
    if not workbook.exists():
        raise FileNotFoundError(workbook)

    wb = load_workbook(workbook)
    if "全市場整合篩選" not in wb.sheetnames:
        raise RuntimeError("duck workbook 缺少「全市場整合篩選」")

    ws = wb["全市場整合篩選"]
    headers = {
        str(ws.cell(row=2, column=c).value or "").strip(): c
        for c in range(1, ws.max_column + 1)
    }
    if "代號" not in headers:
        raise RuntimeError("全市場整合篩選找不到「代號」欄")

    # 不存在的欄位加在最後；現有格式不破壞。
    wanted_cols = [
        "三率狀態","基本面資料狀態","基本面缺漏原因","基本面所需季度",
        "毛利率本季%","毛利率上季%","營業利益率本季%","營業利益率上季%",
        "稅後純益率本季%","稅後純益率上季%","毛差","營差","淨差",
        "三率資料期","三率快照日","三率可得日依據","三率來源",
        "營收狀態","營收資料狀態","營收月增%","營收年增%","累計營收增%",
        "營收資料月","營收快照日","營收可得日依據","營收來源",
        "整合觀察標籤",
    ]
    for name in wanted_cols:
        if name not in headers:
            c = ws.max_column + 1
            ws.cell(row=2, column=c, value=name)
            headers[name] = c

    audit = []
    counts = {
        "stocks": 0, "three_rate_yes": 0, "three_rate_no": 0,
        "three_rate_unknown": 0, "three_rate_na": 0,
        "revenue_yes": 0, "revenue_no": 0, "revenue_unknown": 0,
    }

    for row in range(3, ws.max_row + 1):
        sidv = ws.cell(row=row, column=headers["代號"]).value
        if sidv is None:
            continue
        sid = str(sidv).replace(".0", "").strip()
        if not sid:
            continue
        counts["stocks"] += 1

        ts = three_rate_status(sid, prof_map, vmap)
        rs = revenue_status(sid, rev_map, rev_raw_map)

        if ts == "三率三升": counts["three_rate_yes"] += 1
        elif ts == "未三升": counts["three_rate_no"] += 1
        elif ts == "三率不適用": counts["three_rate_na"] += 1
        else: counts["three_rate_unknown"] += 1

        if rs == "營收三增": counts["revenue_yes"] += 1
        elif rs == "未三增": counts["revenue_no"] += 1
        else: counts["revenue_unknown"] += 1

        ws.cell(row=row, column=headers["三率狀態"], value=ts)
        ws.cell(row=row, column=headers["營收狀態"], value=rs)

        vv = vmap.get(sid)
        if vv is not None:
            ws.cell(row=row, column=headers["基本面資料狀態"], value=vv.get("status"))
            ws.cell(row=row, column=headers["基本面缺漏原因"], value=vv.get("reason"))
            ws.cell(row=row, column=headers["基本面所需季度"], value=vv.get("required"))

        p = prof_map.get(sid, {})
        for k in [
            "毛利率本季%","毛利率上季%","營業利益率本季%","營業利益率上季%",
            "稅後純益率本季%","稅後純益率上季%","毛差","營差","淨差",
            "三率資料期","三率快照日","三率可得日依據","三率來源",
        ]:
            if k in p:
                ws.cell(row=row, column=headers[k], value=p.get(k))

        rraw = rev_map.get(sid) or rev_raw_map.get(sid) or {}
        for k in [
            "營收月增%","營收年增%","累計營收增%","營收資料月",
            "營收快照日","營收可得日依據","營收來源",
        ]:
            if k in rraw:
                ws.cell(row=row, column=headers[k], value=rraw.get(k))

        if rs in {"營收三增","未三增"}:
            ws.cell(row=row, column=headers["營收資料狀態"], value="完整")
        elif sid in rev_raw_map:
            ws.cell(row=row, column=headers["營收資料狀態"], value="成長率不適用/基期不足")
        else:
            ws.cell(row=row, column=headers["營收資料狀態"], value="資料不足")

        # 重新整理標籤：只移除舊的基本面標籤，其餘鴨嘴/估值標籤保留。
        old = str(ws.cell(row=row, column=headers["整合觀察標籤"]).value or "")
        parts = [
            x for x in old.split("｜")
            if x and x not in {"三率三升", "營收三增"}
        ]
        if ts == "三率三升":
            parts.append("三率三升")
        if rs == "營收三增":
            parts.append("營收三增")
        # 去重複但保留順序
        parts = list(dict.fromkeys(parts))
        ws.cell(row=row, column=headers["整合觀察標籤"], value="｜".join(parts))

        name = ""
        if "名稱" in headers:
            name = str(ws.cell(row=row, column=headers["名稱"]).value or "")
        audit.append([
            sid, name, ts, p.get("三率資料期"), rs, rraw.get("營收資料月"),
            vv.get("reason") if vv is not None else None,
        ])

    # 建立／重建全市場驗收工作表，方便直接確認廣達或其他股票。
    if "基本面全市場驗收" in wb.sheetnames:
        del wb["基本面全市場驗收"]
    va = wb.create_sheet("基本面全市場驗收")
    va.append(["代號","名稱","三率狀態","三率資料期","營收狀態","營收資料月","缺漏/說明"])
    for rec in audit:
        va.append(rec)
    va.freeze_panes = "A2"
    va.auto_filter.ref = va.dimensions
    widths = [12,18,22,18,24,16,55]
    for i, width in enumerate(widths, 1):
        va.column_dimensions[chr(64+i)].width = width

    # 驗收摘要
    va["I1"] = "驗收摘要"
    va["I2"] = f"目標日：{target}"
    va["I3"] = f"普通股：{counts['stocks']}"
    va["I4"] = f"三率三升：{counts['three_rate_yes']}"
    va["I5"] = f"三率未三升：{counts['three_rate_no']}"
    va["I6"] = f"三率資料不足：{counts['three_rate_unknown']}"
    va["I7"] = f"三率不適用：{counts['three_rate_na']}"
    va["I8"] = f"營收三增：{counts['revenue_yes']}"
    va["I9"] = f"營收未三增：{counts['revenue_no']}"
    va["I10"] = f"營收資料不足/基期不足：{counts['revenue_unknown']}"

    # 廣達 2382 特別驗收，避免這次問題再悄悄回來。
    q = next((x for x in audit if x[0] == "2382"), None)
    va["I12"] = "2382 廣達驗收"
    va["I13"] = f"三率：{q[2] if q else '未找到'}"
    va["I14"] = f"資料期：{q[3] if q else '—'}"
    va["I15"] = f"營收：{q[4] if q else '未找到'}"
    va["I16"] = f"營收月：{q[5] if q else '—'}"

    wb.save(workbook)
    return counts, q


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--date")
    ap.add_argument("--workbook", default=str(ROOT / "duck_latest.xlsx"))
    ap.add_argument(
        "--repair-missing-core", action="store_true",
        help="一次性補全市場三率判定所需核心季度；日常排程通常不需要開。",
    )
    ap.add_argument("--max-workers", type=int, default=4)
    args = ap.parse_args()

    target = _target_date(args.date)
    workbook = Path(args.workbook)
    if not workbook.is_absolute():
        workbook = (Path.cwd() / workbook).resolve()

    cfg = sd.load_config()
    conn = sqlite3.connect(sd.DB_PATH)
    sd.init_db(conn)

    ids = _all_market_ids(conn, target)
    if not ids:
        raise SystemExit(f"{target} stock.db 找不到全市場普通股")

    print(f"[全市場基本面驗收] target={target} 普通股={len(ids)}", flush=True)

    # 日常先用官方批次最新資料補齊；成本低。
    prof_map, rev_map, rev_raw_map, vmap, source_meta = build_full_market_maps(
        conn, target, ids, cfg
    )

    repair_meta = None
    if args.repair_missing_core:
        repair_meta = repair_missing_core_income(
            conn, target, ids, workers=args.max_workers
        )
        # 補洞後重新建立三率 map / validation。
        prof_map, rev_map, rev_raw_map, vmap, source_meta2 = build_full_market_maps(
            conn, target, ids, cfg
        )
        source_meta["after_repair"] = source_meta2

    counts, quanta = patch_workbook(
        workbook, prof_map, rev_map, rev_raw_map, vmap, target
    )

    report = {
        "target_date": target.isoformat(),
        "workbook": str(workbook),
        "all_market_stock_count": len(ids),
        "counts": counts,
        "source": source_meta,
        "repair": repair_meta,
        "quanta_2382": {
            "stock_id": quanta[0] if quanta else None,
            "name": quanta[1] if quanta else None,
            "three_rate_status": quanta[2] if quanta else None,
            "three_rate_period": quanta[3] if quanta else None,
            "revenue_status": quanta[4] if quanta else None,
            "revenue_month": quanta[5] if quanta else None,
        },
        "rule": {
            "three_rates": "毛利率、營業利益率、稅後淨利率；最新單季 vs 前一單季，三者皆上升",
            "revenue_three_growth": "月增率、年增率、累計營收年增率均 > 設定門檻（目前預設0）",
            "unknown_is_false": False,
        },
    }
    report_path = ROOT / "fundamental_full_market_status.json"
    report_path.write_text(
        json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8"
    )

    print("=" * 70)
    print("[完成] 全市場三率 / 營收驗收已寫回 workbook", flush=True)
    print(json.dumps(report["counts"], ensure_ascii=False), flush=True)
    print("2382廣達：", json.dumps(report["quanta_2382"], ensure_ascii=False), flush=True)
    print(f"報告：{report_path}", flush=True)
    print("=" * 70)
    conn.close()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
